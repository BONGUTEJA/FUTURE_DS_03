"""
FUTURE_DS_03: Marketing Funnel & Conversion Performance Analysis
Module: Excel Executive Workbook Builder for Marketing Analytics
Author: Future Interns - Data Science & Analytics Track
"""

import os
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


def build_funnel_excel_workbook(clean_path: str, output_path: str):
    print(f"[*] Reading cleaned funnel dataset from: {clean_path}")
    df = pd.read_csv(clean_path)

    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    # Styles
    navy_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    light_blue_fill = PatternFill(start_color="DBEAFE", end_color="DBEAFE", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    title_font = Font(name="Calibri", size=16, bold=True, color="1E3A8A")
    thin_border = Border(
        left=Side(style='thin', color='CCCCCC'),
        right=Side(style='thin', color='CCCCCC'),
        top=Side(style='thin', color='CCCCCC'),
        bottom=Side(style='thin', color='CCCCCC')
    )

    # 1. Executive Scorecard Sheet
    ws1 = wb.create_sheet(title="Executive_Scorecard")
    ws1.views.sheetView[0].showGridLines = True
    ws1["A1"] = "MARKETING FUNNEL & CONVERSION PERFORMANCE - EXECUTIVE SCORECARD"
    ws1["A1"].font = title_font

    total_visitors = len(df)
    leads = int(df['Lead_Captured'].sum())
    won = int(df['Closed_Won'].sum())
    spend = float(df['Ad_Spend'].sum())
    rev = float(df['Revenue'].sum())

    kpis = [
        ("Total Website Visitors Audited", total_visitors, "#,##0"),
        ("Leads Captured (Form Submissions)", leads, "#,##0"),
        ("Marketing Qualified Leads (MQL)", int(df['MQL'].sum()), "#,##0"),
        ("Sales Qualified Leads (SQL)", int(df['SQL'].sum()), "#,##0"),
        ("Sales Opportunities / Demos", int(df['Opportunity'].sum()), "#,##0"),
        ("Closed-Won Customer Accounts", won, "#,##0"),
        ("Overall Visitor-to-Customer Conversion Rate", won / total_visitors, "0.00%"),
        ("Lead-to-Customer Conversion Rate", won / leads, "0.00%"),
        ("Total Marketing Investment (Ad Spend)", spend, "$#,##0.00"),
        ("Total Attributed Pipeline Revenue Won", rev, "$#,##0.00"),
        ("Blended Return on Ad Spend (ROAS)", rev / spend, "0.00x"),
        ("Customer Acquisition Cost (CAC)", spend / won, "$#,##0.00"),
        ("Cost per Captured Lead (CPL)", spend / leads, "$#,##0.00")
    ]

    ws1.cell(row=3, column=1, value="Core Growth & Conversion KPI").fill = navy_fill
    ws1.cell(row=3, column=1).font = header_font
    ws1.cell(row=3, column=2, value="Realized Metric").fill = navy_fill
    ws1.cell(row=3, column=2).font = header_font

    for i, (kpi, val, fmt) in enumerate(kpis, start=4):
        c1 = ws1.cell(row=i, column=1, value=kpi)
        c2 = ws1.cell(row=i, column=2, value=val)
        c1.border = thin_border
        c2.border = thin_border
        c2.number_format = fmt
        c2.alignment = Alignment(horizontal="right")
        if i % 2 == 1:
            c1.fill = light_blue_fill

    # 2. Funnel Progression Sheet
    ws2 = wb.create_sheet(title="Funnel_Progression")
    ws2.views.sheetView[0].showGridLines = True
    ws2["A1"] = "STAGE-BY-STAGE FUNNEL CONVERSION & DROP-OFF AUDIT"
    ws2["A1"].font = title_font

    stages = [
        ("1. Website Visitors", total_visitors),
        ("2. Leads Captured", leads),
        ("3. Marketing Qualified Leads (MQL)", int(df['MQL'].sum())),
        ("4. Sales Qualified Leads (SQL)", int(df['SQL'].sum())),
        ("5. Sales Opportunities", int(df['Opportunity'].sum())),
        ("6. Closed-Won Customers", won)
    ]
    f_headers = ["Funnel Stage", "Prospect Volume", "Share of Total", "Step Conversion Rate", "Step Drop-Off Rate"]
    for col_idx, h in enumerate(f_headers, 1):
        c = ws2.cell(row=3, column=col_idx, value=h)
        c.fill = navy_fill
        c.font = header_font

    for r_idx, (st_name, st_count) in enumerate(stages, start=4):
        ws2.cell(row=r_idx, column=1, value=st_name).border = thin_border
        c2 = ws2.cell(row=r_idx, column=2, value=st_count)
        c2.border = thin_border
        c2.number_format = "#,##0"
        c3 = ws2.cell(row=r_idx, column=3, value=st_count / total_visitors)
        c3.border = thin_border
        c3.number_format = "0.0%"

        if r_idx == 4:
            step_c = 1.0
            drop_o = 0.0
        else:
            prev_cnt = stages[r_idx - 5][1]
            step_c = st_count / prev_cnt
            drop_o = 1.0 - step_c

        c4 = ws2.cell(row=r_idx, column=4, value=step_c)
        c4.border = thin_border
        c4.number_format = "0.0%"

        c5 = ws2.cell(row=r_idx, column=5, value=drop_o)
        c5.border = thin_border
        c5.number_format = "0.0%"

    # 3. Channel Performance Sheet
    ws3 = wb.create_sheet(title="Channel_Performance")
    ws3.views.sheetView[0].showGridLines = True
    ws3["A1"] = "MARKETING CHANNEL EFFICIENCY (CAC & ROAS)"
    ws3["A1"].font = title_font

    ch_df = df.groupby('Channel').agg(
        Visitors=('Lead_ID', 'count'),
        Leads=('Lead_Captured', 'sum'),
        Won=('Closed_Won', 'sum'),
        Spend=('Ad_Spend', 'sum'),
        Revenue=('Revenue', 'sum')
    ).reset_index()

    ch_headers = ["Acquisition Channel", "Visitors", "Leads", "Won Deals", "Conv Rate", "Ad Spend ($)", "Revenue ($)", "CAC ($)", "ROAS"]
    for col_idx, h in enumerate(ch_headers, 1):
        c = ws3.cell(row=3, column=col_idx, value=h)
        c.fill = navy_fill
        c.font = header_font

    for r_idx, row in enumerate(ch_df.itertuples(), start=4):
        ws3.cell(row=r_idx, column=1, value=row.Channel).border = thin_border
        ws3.cell(row=r_idx, column=2, value=row.Visitors).border = thin_border
        ws3.cell(row=r_idx, column=3, value=row.Leads).border = thin_border
        ws3.cell(row=r_idx, column=4, value=row.Won).border = thin_border

        c5 = ws3.cell(row=r_idx, column=5, value=row.Won / row.Visitors)
        c5.border = thin_border
        c5.number_format = "0.00%"

        c6 = ws3.cell(row=r_idx, column=6, value=row.Spend)
        c6.border = thin_border
        c6.number_format = "$#,##0.00"

        c7 = ws3.cell(row=r_idx, column=7, value=row.Revenue)
        c7.border = thin_border
        c7.number_format = "$#,##0.00"

        c8 = ws3.cell(row=r_idx, column=8, value=row.Spend / row.Won if row.Won > 0 else 0)
        c8.border = thin_border
        c8.number_format = "$#,##0.00"

        c9 = ws3.cell(row=r_idx, column=9, value=row.Revenue / row.Spend if row.Spend > 0 else 0)
        c9.border = thin_border
        c9.number_format = "0.00x"

    # 4. Cleaned Lead Data Sample (Top 500 rows)
    ws4 = wb.create_sheet(title="Lead_Data_Sample")
    ws4.views.sheetView[0].showGridLines = True
    sample_df = df.head(500)
    for c_idx, col_name in enumerate(sample_df.columns, 1):
        cell = ws4.cell(row=1, column=c_idx, value=col_name)
        cell.fill = navy_fill
        cell.font = header_font

    for r_idx, row in enumerate(sample_df.values, 2):
        for c_idx, val in enumerate(row, 1):
            ws4.cell(row=r_idx, column=c_idx, value=val).border = thin_border

    # Auto column width
    for sheet in wb.worksheets:
        for col in sheet.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = get_column_letter(col[0].column)
            sheet.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 40)

    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    wb.save(output_path)
    print(f"[+] Marketing funnel Excel workbook created at: {output_path}")


if __name__ == '__main__':
    base_dir = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    clean_csv = os.path.join(base_dir, 'data', 'cleaned_marketing_funnel_data.csv')
    wb_out = os.path.join(base_dir, 'dashboards', 'marketing_funnel_workbook.xlsx')
    build_funnel_excel_workbook(clean_csv, wb_out)
